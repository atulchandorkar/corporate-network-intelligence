import pandas as pd, openpyxl
from openpyxl.styles import Font, PatternFill
from pathlib import Path
ROOT = Path(__file__).resolve().parents[2]
OUT = str(ROOT / "01_data" / "synthetic") + "/"   # written by gen.py
df = pd.read_pickle(OUT + 'df.pkl'); E = pd.read_pickle(OUT + 'E.pkl')
csv = df.copy()
for c in ["BOOKING_DATE", "VINTAGE", "BATCH_DATE"]:
    csv[c] = pd.to_datetime(csv[c]).dt.strftime("%Y-%m-%d")
csv.to_csv(OUT + "qib_synthetic_transactions.csv", index=False, na_rep="NULL")

BANKS = [(4,"QISB","Qatar Islamic Bank"),(1,"QNBA","Qatar National Bank"),(2,"CBQA","Commercial Bank of Qatar"),(3,"DOHB","Doha Bank"),
         (5,"ABQQ","Ahli Bank"),(6,"QIIB","Qatar International Islamic Bank"),(7,"MAFR","Masraf Al Rayan"),(8,"BRWA","Dukhan Bank"),
         (9,"BBME","HSBC Bank Middle East - Qatar"),(10,"SCBL","Standard Chartered Bank - Qatar")]
banks = pd.DataFrame(BANKS, columns=["BANK_CD","IBAN_BANK_ID","BANK_NAME"])
ents = E[["name","kind","industry","bank","is_cust","cif","cp_key"]].rename(columns={
    "name":"ENTITY_NAME","kind":"ENTITY_TYPE","industry":"INDUSTRY","bank":"HOME_BANK","is_cust":"IS_QIB_CUSTOMER","cif":"CIF_ID","cp_key":"UNIQUE_ID"})
used = set(df.CUSTOMER_NAME) | set(df.CP_CUSTOMER_NAME)
ents = ents[ents.ENTITY_NAME.isin(used)]
dd = pd.read_excel(ROOT / "08_documentation" / "data_dictionary" / "qatar_bank_txn_data_structure.xlsx", sheet_name="Data_Dictionary")

xl = df.copy()
for c in ["BOOKING_DATE","VINTAGE","BATCH_DATE"]: xl[c] = pd.to_datetime(xl[c])
with pd.ExcelWriter(OUT + "qib_synthetic_transactions.xlsx", engine="openpyxl", date_format="m/d/yyyy", datetime_format="m/d/yyyy") as w:
    xl.to_excel(w, sheet_name="Transactions", index=False, na_rep="NULL")
    ents.to_excel(w, sheet_name="Entities", index=False)
    banks.to_excel(w, sheet_name="Bank_Codes", index=False)
    dd.to_excel(w, sheet_name="Data_Dictionary", index=False)
    for ws in w.book.worksheets:
        for c in ws[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="7A1E3A")
        ws.freeze_panes = "A2"
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(45, max(10, max(len(str(c.value or "")) for c in col[:200]) + 2))
    w.book["Transactions"].auto_filter.ref = w.book["Transactions"].dimensions
print(ents.ENTITY_TYPE.value_counts().to_dict())
